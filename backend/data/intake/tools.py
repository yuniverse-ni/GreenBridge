"""Read-only intake tools. Reuse parser/service contracts for authoritative data."""

import calendar
import csv
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
from collections import defaultdict

from ..mapping import EXPECTED_UNITS, canonical_unit, header_key, map_column
from ..models import ActivityData, BASIC_FINANCIAL_FIELDS, ConfirmedMapping, FinancialData, NormalizedDataset, ParsedDocument
from ..parser import DataImportError, parse_upload
from .models import DocumentClassification, EvidenceAssessment, EvidenceGap, Inspection, Readiness, Upload


def inspect_file(upload: Upload) -> Inspection:
    kind = Path(upload.filename).suffix.lower().lstrip(".")
    inspection = Inspection(file_id=upload.file_id, filename=upload.filename,
                            source_type=kind if kind in {"csv", "xlsx"} else "unsupported",
                            sha256=sha256(upload.content).hexdigest())
    try:
        if kind == "csv":
            text = upload.content.decode("utf-8-sig")
            first = text.splitlines()[0] if text.splitlines() else ""
            delimiter = ";" if ";" in first and "," not in first else "\t" if "\t" in first and "," not in first else ","
            tables = [list(csv.reader(StringIO(text), delimiter=delimiter, strict=True))]
        elif kind == "xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(upload.content), read_only=True, data_only=False)
            try:
                tables = [list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets]
            finally:
                workbook.close()
        else:
            raise ValueError("Only CSV/XLSX intake is supported; no OCR is configured")
        for rows in tables:
            if not rows:
                continue
            headers = [str(cell) if cell is not None else "" for cell in rows[0]]
            inspection.headers.extend(headers)
            keys = [header_key(header) for header in headers]
            metric = next((keys.index(key) for key in ("metric", "field") if key in keys), None)
            if metric is not None and "value" in keys:
                unit_index = keys.index("unit") if "unit" in keys else None
                for row in rows[1:]:
                    if metric < len(row) and row[metric] is not None:
                        label = str(row[metric])
                        inspection.metric_labels.append(label)
                        unit = row[unit_index] if unit_index is not None and unit_index < len(row) else None
                        mapping, _ = map_column(label, unit)
                        if mapping.mapped_field:
                            inspection.supported_fields.append(mapping.mapped_field)
            else:
                for label in headers:
                    mapping, _ = map_column(label)
                    if mapping.mapped_field:
                        inspection.supported_fields.append(mapping.mapped_field)
        inspection.supported_fields = sorted(set(inspection.supported_fields))
    except Exception:
        inspection.error = "Unable to inspect structured file; provide a valid CSV/XLSX export."
    return inspection


def classify_document(inspection: Inspection) -> DocumentClassification:
    fields = set(inspection.supported_fields)
    categories = []
    if fields & set(FinancialData.model_fields):
        categories.append("financial_records")
    if "electricity_kwh" in fields:
        categories.append("electricity_evidence")
    if fields & {"diesel_litres", "petrol_litres", "natural_gas_m3", "lpg_kg"}:
        categories.append("fuel_records")
    if not categories and fields & {"employees", "floor_area_m2"}:
        categories.append("operational_records")
    kind = "mixed" if len(categories) > 1 else categories[0] if categories else "unknown"
    confidence = "HIGH" if categories else "UNKNOWN"
    reasoning = "Recognized structured field headers; classification does not verify evidence." if categories else "No reliable document category identified."
    if not categories and not inspection.error:
        stem = header_key(Path(inspection.filename).stem)
        hints = [("tnb", "electricity_evidence"), ("electricity", "electricity_evidence"),
                 ("diesel", "fuel_records"), ("petrol", "fuel_records"),
                 ("accounts", "financial_records"), ("financial", "financial_records"),
                 ("supplier", "supplier_evidence")]
        matches = {category for hint, category in hints if hint in stem}
        if len(matches) == 1:
            kind, confidence = matches.pop(), "LOW"
            reasoning = "Tentative filename hint only; content interpretation still needs validation."
    strategy = f"deterministic_{inspection.source_type}" if inspection.source_type in {"csv", "xlsx"} else "unsupported"
    return DocumentClassification(document_type=kind, confidence=confidence, reasoning=reasoning, extraction_strategy=strategy)


def parse_csv(upload: Upload, confirmations: list[ConfirmedMapping]) -> ParsedDocument:
    if Path(upload.filename).suffix.lower() != ".csv":
        raise DataImportError("CSV tool requires a .csv file")
    return parse_upload(upload.content, upload.filename, confirmed_mappings=confirmations)


def parse_excel(upload: Upload, confirmations: list[ConfirmedMapping]) -> ParsedDocument:
    if Path(upload.filename).suffix.lower() != ".xlsx":
        raise DataImportError("Excel tool requires a .xlsx file")
    return parse_upload(upload.content, upload.filename, confirmed_mappings=confirmations)


def validate_units(mapping: ConfirmedMapping) -> None:
    expected = EXPECTED_UNITS.get(mapping.mapped_field)
    actual = canonical_unit(mapping.unit)
    if expected and not actual:
        raise ValueError(f"Explicit input unit is required for {mapping.mapped_field}")
    if actual != expected and not (expected == "kWh" and actual in {"Wh", "MWh"}):
        raise ValueError(f"Cannot normalize {actual} to {expected}")


def validate_financial_data(dataset: NormalizedDataset) -> None:
    FinancialData.model_validate(dataset.financial.model_dump())


def detect_duplicate_evidence(inspections: list[Inspection], documents: list[ParsedDocument]) -> list[EvidenceGap]:
    gaps = []
    digests = defaultdict(list)
    for item in inspections:
        digests[item.sha256].append(item.filename)
    for names in digests.values():
        if len(names) > 1:
            gaps.append(EvidenceGap(gap_type="DUPLICATE_FILE", severity="ERROR",
                                    message=f"Identical evidence was uploaded more than once: {', '.join(names)}.",
                                    recommended_action="Remove duplicate uploads before continuing."))
    fields = defaultdict(list)
    for document in documents:
        for item in document.values:
            if item.field in {"electricity_kwh", "diesel_litres", "petrol_litres", "natural_gas_m3", "lpg_kg"}:
                fields[item.field].append(item.provenance.reporting_period)
    for field, periods in fields.items():
        if len(periods) > 1 and (None in periods or len(set(periods)) != len(periods)
                                or any(len(period or "") != 7 for period in periods)):
            gaps.append(EvidenceGap(gap_type="OVERLAPPING_EVIDENCE", affected_field=field, severity="ERROR",
                                    message="Activity evidence overlaps or cannot be proven to cover distinct months.",
                                    recommended_action="Remove duplicates or provide one non-overlapping monthly/annual export."))
    return gaps


def check_reporting_period_coverage(dataset: NormalizedDataset) -> list[EvidenceGap]:
    gaps = []
    year = dataset.metadata.reporting_year
    for field, value in dataset.activity.model_dump().items():
        if value is None:
            continue
        periods = [item.reporting_period for item in dataset.provenance.get(field, [])]
        if not periods or any(period is None for period in periods):
            gaps.append(EvidenceGap(gap_type="MISSING_PERIOD", affected_field=field, severity="WARNING",
                                    message="Activity evidence has no reporting period.",
                                    recommended_action="Supply dated evidence or explicitly continue with incomplete evidence."))
        elif any(len(period) == 4 for period in periods):
            continue  # Explicit annual aggregate, not an assertion of independent verification.
        elif year is not None:
            # A single dated transaction is not proof of a complete month.
            months = {period for period in periods if len(period) == 7}
            missing = [f"{year}-{month:02d}" for month in range(1, 13) if f"{year}-{month:02d}" not in months]
            if missing:
                names = ", ".join(calendar.month_name[int(period[-2:])] for period in missing)
                gaps.append(EvidenceGap(gap_type="INCOMPLETE_PERIOD", affected_field=field, severity="WARNING",
                                        message=f"Annual monthly coverage is incomplete for {year}; missing {names}.",
                                        missing_periods=missing,
                                        recommended_action="Upload remaining monthly records or explicitly continue with incomplete evidence."))
    return gaps


def assess_evidence_confidence(uploads: list[Upload], documents: list[ParsedDocument]) -> list[EvidenceAssessment]:
    assessments = []
    for upload, document in zip(uploads, documents):
        for item in document.values:
            source = item.provenance
            method = source.mapping_method
            status = "AI_MAPPED_CONFIRMED" if method == "ai_mapped_confirmed" else "HUMAN_CONFIRMED" if method == "human_confirmed" else "DETERMINISTIC_ACCEPTED"
            assessments.append(EvidenceAssessment(
                file_id=upload.file_id, field=item.field, source_file=upload.filename, row=source.row,
                column=source.original_metric or source.column, sheet=source.sheet,
                mapping_status=status, mapping_confidence=source.mapping_confidence,
                semantic_confidence=source.semantic_confidence,
                evidence_quality=upload.evidence_quality, quality_basis=upload.quality_basis,
                confirmation_status=source.confirmation_status,
            ))
    return assessments


def assess_readiness(dataset: NormalizedDataset | None, gaps: list[EvidenceGap], *, blocked: bool = False) -> dict[str, Readiness]:
    """Presence/coverage screening only, never a conclusion about actual eligibility."""
    if dataset is None or blocked:
        return {module: Readiness(status="BLOCKED", reasons=["Resolve validation, ambiguity or duplicate-evidence issues first."])
                for module in ("carbon_analysis", "financial_analysis", "transition_analysis")}
    metadata = dataset.metadata
    context_missing = [field for field in ("company_id", "reporting_year") if getattr(metadata, field) is None]
    carbon_fields = [field for field, value in dataset.activity.model_dump().items() if value is not None]
    finance_fields = [field for field, value in dataset.financial.model_dump().items() if value is not None]
    carbon_gaps = [gap.message for gap in gaps if gap.affected_field in ActivityData.model_fields]
    financial_missing = [field for field in BASIC_FINANCIAL_FIELDS if getattr(dataset.financial, field) is None]
    carbon_reasons = ([f"Missing context: {', '.join(context_missing)}."] if context_missing else []) + carbon_gaps
    context_quality_uncertain = any(
        source.evidence_quality in {"UNKNOWN", "ESTIMATED_PROXY"}
        for field in ("company_id", "reporting_year") for source in dataset.provenance.get(field, [])
    )
    if context_quality_uncertain:
        carbon_reasons.append("Company/reporting-year source quality is unknown or estimated.")
    if not carbon_fields:
        carbon_reasons.append("No energy or fuel evidence is available.")
    # Missing fuel sources are not assumed zero; state this rather than claiming full baseline readiness.
    if carbon_fields and len(carbon_fields) < len(ActivityData.model_fields):
        carbon_reasons.append("Some activity categories are missing; coverage/applicability must be reviewed.")
    for field in carbon_fields:
        if any(item.evidence_quality in {"UNKNOWN", "ESTIMATED_PROXY"} for item in dataset.provenance.get(field, [])):
            carbon_reasons.append(f"Source quality for {field} is unknown or estimated; not externally verified.")
    carbon_status = "BLOCKED" if context_missing or not carbon_fields else "PARTIAL" if carbon_reasons else "READY"
    finance_reasons = ([f"Missing context: {', '.join(context_missing)}."] if context_missing else [])
    if context_quality_uncertain:
        finance_reasons.append("Company/reporting-year source quality is unknown or estimated.")
    if financial_missing:
        finance_reasons.append(f"Missing financial fields: {', '.join(financial_missing)}.")
    for field in finance_fields:
        if any(item.evidence_quality in {"UNKNOWN", "ESTIMATED_PROXY"} for item in dataset.provenance.get(field, [])):
            finance_reasons.append(f"Source quality for {field} is unknown or estimated; not externally verified.")
    finance_status = "BLOCKED" if context_missing or not finance_fields else "PARTIAL" if finance_reasons else "READY"
    transition_status = "READY" if carbon_status == finance_status == "READY" else "BLOCKED" if "BLOCKED" in {carbon_status, finance_status} else "PARTIAL"
    return {
        "carbon_analysis": Readiness(status=carbon_status, reasons=carbon_reasons or ["Required structured activity/context fields and annual coverage are present; evidence is not externally verified."]),
        "financial_analysis": Readiness(status=finance_status, reasons=finance_reasons or ["Required financial/context fields are present; no financial analysis has been performed."]),
        "transition_analysis": Readiness(status=transition_status, reasons=["Depends on carbon and financial data readiness; no transition calculations have been performed."]),
    }


def find_evidence_gaps(dataset: NormalizedDataset, modules: list[str]) -> list[EvidenceGap]:
    gaps = []
    for field in ("company_id", "reporting_year"):
        if getattr(dataset.metadata, field) is None:
            gaps.append(EvidenceGap(gap_type="MISSING_CONTEXT", affected_field=field, severity="ERROR",
                                    message=f"{field} is missing.", recommended_action=f"Upload a record containing {field}."))
    if set(modules) & {"carbon_analysis", "transition_analysis"}:
        if all(value is None for value in dataset.activity.model_dump().values()):
            gaps.append(EvidenceGap(gap_type="MISSING_ACTIVITY", severity="ERROR",
                                    message="Carbon analysis was requested but no energy/fuel evidence exists.",
                                    recommended_action="Upload electricity or fuel records with explicit units and periods."))
        else:
            missing = [field for field, value in dataset.activity.model_dump().items() if value is None]
            if missing:
                gaps.append(EvidenceGap(gap_type="ACTIVITY_APPLICABILITY", severity="WARNING",
                                        message=f"Missing activity categories are not assumed zero: {', '.join(missing)}.",
                                        recommended_action="Supply applicable activity evidence or explicitly continue with a partial baseline."))
        gaps.extend(check_reporting_period_coverage(dataset))
    if set(modules) & {"financial_analysis", "transition_analysis"}:
        for field in BASIC_FINANCIAL_FIELDS:
            if getattr(dataset.financial, field) is None:
                gaps.append(EvidenceGap(gap_type="MISSING_FINANCIAL", affected_field=field, severity="WARNING",
                                        message=f"{field} is missing; available financial values do not imply zero.",
                                        recommended_action=f"Upload financial evidence containing {field}, or explicitly continue incomplete."))
    relevant_fields = {"company_id", "reporting_year"}
    if set(modules) & {"carbon_analysis", "transition_analysis"}:
        relevant_fields.update(ActivityData.model_fields)
    if set(modules) & {"financial_analysis", "transition_analysis"}:
        relevant_fields.update(FinancialData.model_fields)
    relevant_sources = [source for field in relevant_fields for source in dataset.provenance.get(field, [])]
    if any(source.evidence_quality == "UNKNOWN" for source in relevant_sources):
        gaps.append(EvidenceGap(gap_type="UNKNOWN_EVIDENCE_QUALITY", severity="WARNING",
                                message="Evidence quality is undeclared; matching confidence does not establish source reliability.",
                                recommended_action="Review source evidence quality or explicitly continue with unverified evidence."))
    if any(source.evidence_quality == "ESTIMATED_PROXY" for source in relevant_sources):
        gaps.append(EvidenceGap(gap_type="ESTIMATED_EVIDENCE", severity="WARNING",
                                message="Some records are declared estimates/proxies, not measured evidence.",
                                recommended_action="Provide measured/invoice records or explicitly continue with estimated evidence."))
    return gaps
