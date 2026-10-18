"""Inspect CSV/XLSX and normalize deterministically, with mapping diagnostics.

Partial imports return unmapped fields and warnings. Files with no usable data,
invalid values, duplicate aliases or malformed structure raise DataImportError.
"""

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
import re
from typing import Any, Iterable

from .mapping import EXPECTED_UNITS, MEASUREMENT_PATTERN, canonical_unit, header_key, map_column, numeric_value, UnitInterpretationError
from .models import ConfirmedMapping, FIELD_GROUPS, GROUP_MODELS, FieldMapping, MappingLocation, MappingWarning, ParsedDocument, ParsedValue, Provenance


class DataImportError(ValueError):
    """Controlled failure; ``as_dict()`` exposes JSON-ready inspection diagnostics."""

    def __init__(self, message: str, *, document: ParsedDocument | None = None):
        super().__init__(message)
        self.document = document
        self.unmapped_fields = document.unmapped_fields if document else []
        self.warnings = document.warnings if document else []
        self.mappings = document.mappings if document else []

    def as_dict(self) -> dict:
        return {"error": str(self), "unmapped_fields": self.unmapped_fields,
                "warnings": [warning.model_dump() for warning in self.warnings],
                "mappings": [mapping.model_dump() for mapping in self.mappings]}


MISSING = {"", "null", "none", "n/a", "na"}
CONTROL_KEYS = {"field", "metric", "value", "unit", "period", "reportingperiod", "date", "confidence"}


def _missing(value: Any) -> bool:
    return value is None or isinstance(value, str) and value.strip().lower() in MISSING


def _value_unit(value: Any) -> str | None:
    """Inspect explicit cell units, never a numeric magnitude."""
    if not isinstance(value, str):
        return None
    currency = re.match(r"^\(?\s*(RM|MYR|USD|SGD|EUR|[$€£])\s*", value.strip(), re.IGNORECASE)
    if currency:
        return currency.group(1)
    suffix = re.fullmatch(MEASUREMENT_PATTERN, value.strip())
    return suffix.group(2) if suffix else None


def _record_mapping(document: ParsedDocument, mapping: FieldMapping, message: str | None,
                    row: int | None, sheet: str | None) -> None:
    document.mapping_locations.append(MappingLocation(mapping_index=len(document.mappings), sheet=sheet))
    document.mappings.append(mapping.model_copy(update={"row": row}))
    if not mapping.mapped_field:
        if mapping.original_column not in document.unmapped_fields:
            document.unmapped_fields.append(mapping.original_column)
        document.warnings.append(MappingWarning(
            code="AMBIGUOUS_MAPPING" if mapping.mapping_method == "ambiguous" else "UNMAPPED_FIELD",
            message=message or "Field not mapped", source_file=document.source_file,
            original_column=mapping.original_column, row=row, sheet=sheet,
        ))


def _parse_rows(rows: Iterable[tuple], document: ParsedDocument, kind: str, sheet: str | None,
                confirmations: list[ConfirmedMapping]) -> None:
    iterator = iter(rows)
    headers = list(next(iterator, ()))
    while headers and _missing(headers[-1]):
        headers.pop()
    keys = [header_key(header) for header in headers]
    if not keys:
        return
    filename = document.source_file
    if any(not key for key in keys) or len(set(keys)) != len(keys):
        raise DataImportError(f"{filename}/{sheet or 'CSV'}: empty or duplicate column names", document=document)
    metric_keys = [key for key in ("field", "metric") if key in keys]
    long_format = bool(metric_keys) and "value" in keys
    if len(metric_keys) > 1 or sum(key in keys for key in ("period", "reportingperiod", "date")) > 1:
        raise DataImportError(f"{filename}: conflicting metric or period columns", document=document)
    header_mappings = {}
    for key, header in zip(keys, headers):
        if key in CONTROL_KEYS:
            continue
        mapping, message = map_column(str(header))
        header_mappings[key] = (mapping, message)
        if long_format:
            _record_mapping(document, mapping, message, None, sheet)
    targets = [mapping.mapped_field for mapping, _ in header_mappings.values() if mapping.mapped_field]
    if len(set(targets)) != len(targets):
        for mapping, message in header_mappings.values():
            _record_mapping(document, mapping, message, None, sheet)
        raise DataImportError(f"{filename}: multiple columns map to the same normalized field", document=document)
    for row_number, raw in enumerate(iterator, 2):
        if all(_missing(cell) for cell in raw):
            continue
        location = f"{filename}/{sheet or 'CSV'} row {row_number}"
        if len(raw) > len(keys) and any(not _missing(cell) for cell in raw[len(keys):]):
            raise DataImportError(f"{location}: more values than headers", document=document)
        cells = dict(zip(keys, list(raw) + [None] * max(0, len(keys) - len(raw))))
        period = next((cells[key] for key in ("period", "reportingperiod", "date") if key in cells), None)
        if isinstance(period, (date, datetime)):
            period = period.date().isoformat() if isinstance(period, datetime) else period.isoformat()
        period = None if _missing(period) else str(period).strip()
        confidence = cells.get("confidence")
        confidence = None if _missing(confidence) else confidence
        if long_format:
            label = str(cells[metric_keys[0]])
            entries = [(label, cells["value"], str(headers[keys.index("value")]), cells.get("unit"))]
        else:
            entries = [(str(header), cells[key], str(header), None)
                       for header, key in zip(headers, keys) if key not in CONTROL_KEYS]
        row_targets = set()
        for label, original_value, column, explicit_unit in entries:
            mapping, message = map_column(label, None if _missing(explicit_unit) else explicit_unit)
            matches = [decision for decision in confirmations if decision.original_column == label
                       and decision.sheet == sheet and decision.row in (None, row_number)]
            if len(matches) > 1:
                raise DataImportError(f"{location}: conflicting human mapping decisions", document=document)
            decision = matches[0] if matches else None
            if decision:
                expected = EXPECTED_UNITS.get(decision.mapped_field)
                unit = canonical_unit(decision.unit)
                if expected and not unit:
                    raise DataImportError(f"{location}: human confirmation must specify input unit ({expected})", document=document)
                observed = canonical_unit(explicit_unit) if not _missing(explicit_unit) else mapping.detected_unit
                # Preserve explicit units in unfamiliar headers too; a human mapping
                # cannot silently rewrite the measurement's declared source unit.
                declared = re.findall(r"(?i)\b(MWh|kWh|Wh|litres?|liters?|L|kg|m[23]|RM|MYR|USD|SGD|EUR)\b", label.replace("_", " ").replace("²", "2").replace("³", "3"))
                observed_units = [canonical_unit(item) for item in declared] + ([observed] if observed else [])
                if any(item != unit for item in observed_units):
                    raise DataImportError(f"{location}: confirmation contradicts original declared unit", document=document)
                mapping = FieldMapping(original_column=label, mapped_field=decision.mapped_field,
                                       detected_unit=unit, normalized_unit=expected, mapping_confidence=None,
                                       mapping_method="ai_mapped_confirmed" if decision.ai_suggested else "human_confirmed")
                message = None
            cell_unit = None if mapping.mapped_field in ("company_id", "company_name") else _value_unit(original_value)
            if cell_unit and mapping.detected_unit is None:
                mapping, message = map_column(label, cell_unit)
            if not mapping.mapped_field:
                _record_mapping(document, mapping, message, row_number, sheet)
                continue
            if mapping.mapped_field in row_targets:
                raise DataImportError(f"{location}: duplicate aliases for {mapping.mapped_field}", document=document)
            row_targets.add(mapping.mapped_field)
            if _missing(original_value):
                continue
            try:
                field = mapping.mapped_field
                if field in ("company_id", "company_name"):
                    value, detected = original_value, mapping.detected_unit
                    if isinstance(value, str) and value.startswith("="):
                        raise ValueError("Formula cells are unsupported; export values first")
                else:
                    value, detected = numeric_value(original_value, mapping)
                mapping = mapping.model_copy(update={"detected_unit": detected, "row": row_number})
                validated = GROUP_MODELS[FIELD_GROUPS[field]].model_validate({field: value})
                provenance = Provenance(
                    source_file=filename, source_type=kind, reporting_period=period,
                    extraction_method=f"deterministic_{kind}", confidence=confidence,
                    row=row_number, column=column, sheet=sheet, original_value=original_value,
                    original_metric=label if long_format else None,
                    detected_unit=detected, normalized_unit=mapping.normalized_unit,
                    mapping_method=mapping.mapping_method, mapping_confidence=mapping.mapping_confidence,
                    confirmation_status="CONFIRMED" if decision else "ACCEPTED",
                    confirmed_by=decision.actor if decision else None,
                    confirmed_at=decision.decided_at if decision else None,
                    semantic_confidence=decision.semantic_confidence if decision else None,
                )
                document.values.append(ParsedValue(field=field, value=getattr(validated, field), provenance=provenance))
                _record_mapping(document, mapping, None, row_number, sheet)
            except UnitInterpretationError as exc:
                rejected = mapping.model_copy(update={"mapped_field": None, "mapping_method": "ambiguous"})
                _record_mapping(document, rejected, str(exc), row_number, sheet)
            except (ValueError, TypeError) as exc:
                raise DataImportError(f"{location}, {column}: {exc}", document=document) from exc


def parse_upload(content: bytes, filename: str, *, confirmed_mappings: list[ConfirmedMapping] | None = None) -> ParsedDocument:
    """Return validated values plus mapping diagnostics, without persistence."""
    kind = Path(filename).suffix.lower().lstrip(".")
    document = ParsedDocument(source_file=filename)
    confirmations = [ConfirmedMapping.model_validate(item) for item in confirmed_mappings or []]
    if kind not in {"csv", "xlsx"}:
        raise DataImportError("Only CSV and XLSX inputs are supported", document=document)
    try:
        if kind == "csv":
            decoded = content.decode("utf-8-sig")
            header = decoded.splitlines()[0] if decoded.splitlines() else ""
            delimiter = ";" if ";" in header and "," not in header else "\t" if "\t" in header and "," not in header else ","
            _parse_rows(csv.reader(StringIO(decoded), delimiter=delimiter, strict=True), document, kind, None, confirmations)
        else:
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
            try:
                for sheet in workbook.worksheets:
                    _parse_rows(sheet.iter_rows(values_only=True), document, kind, sheet.title, confirmations)
            finally:
                workbook.close()
    except DataImportError:
        raise
    except Exception as exc:
        raise DataImportError(f"{filename}: unable to parse {kind}: {exc}", document=document) from exc
    if not document.values:
        raise DataImportError(f"{filename}: no populated supported fields; review mapping diagnostics", document=document)
    return document


def parse_file(path: str | Path) -> ParsedDocument:
    source = Path(path)
    return parse_upload(source.read_bytes(), source.name)
